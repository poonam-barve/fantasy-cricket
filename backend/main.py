import os
import threading
import time
from datetime import datetime, timedelta
try:
    from dotenv import load_dotenv
except ImportError:
    load_dotenv = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

if load_dotenv:
    load_dotenv()
    # Let a local override file win during development without affecting deploys.
    load_dotenv(dotenv_path=os.path.join(os.path.dirname(os.path.dirname(__file__)), ".env.local"), override=True)

from fastapi import FastAPI, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse

from backend.config import IST, get_current_datetime, get_current_date_key, is_current_datetime_overridden
from backend.database import init_db, get_db
from backend.firebase_setup import init_firebase
from backend.services import data_service
from backend.services.scraper import compute_toss_time
from backend.services.venue_stats import prime_today_venue_cache
from backend.services.weekend_tournament_service import prime_weekend_tournament_cache
from backend.services import super_team_service
from backend.models.tournament import Tournament
from backend.routes import auth, matches, players, teams, scores, leaderboard, admin, weekend_tournament, super_team, achievements

app = FastAPI(title="Fantasy Cricket API")

# CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://localhost:3000",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Register API routes
app.include_router(auth.router)
app.include_router(matches.router)
app.include_router(players.router)
app.include_router(teams.router)
app.include_router(scores.router)
app.include_router(leaderboard.router)
app.include_router(admin.router)
app.include_router(weekend_tournament.router)
app.include_router(super_team.router)
app.include_router(super_team.admin_router)
app.include_router(achievements.router)


@app.get("/api/runtime/current-time")
async def runtime_current_time():
    now = get_current_datetime()
    return {
        "iso": now.isoformat(),
        "date": now.strftime("%Y-%m-%d"),
        "time": now.strftime("%H:%M"),
        "timezone": "Asia/Kolkata",
        "overridden": is_current_datetime_overridden(),
    }

# Tournament singleton
tournament = Tournament()
bootstrap_lock = threading.Lock()
bootstrap_started = False
bootstrap_error = None
bootstrap_ready = False
bootstrap_warmup_complete = False
bootstrap_warmup_error = None


def seed_db_if_needed():
    if str(os.environ.get("SKIP_DB_SEED", "")).strip().lower() in {"1", "true", "yes"}:
        print("Skipping database seed because SKIP_DB_SEED is enabled")
        return

    db = get_db()
    counts = {
        "players": db.execute("SELECT COUNT(*) as cnt FROM players").fetchone(),
        "matches": db.execute("SELECT COUNT(*) as cnt FROM matches").fetchone(),
        "users": db.execute("SELECT COUNT(*) as cnt FROM users").fetchone(),
    }
    player_count = counts["players"]["cnt"] if isinstance(counts["players"], dict) else counts["players"][0]
    match_count = counts["matches"]["cnt"] if isinstance(counts["matches"], dict) else counts["matches"][0]
    user_count = counts["users"]["cnt"] if isinstance(counts["users"], dict) else counts["users"][0]

    if player_count > 0 and match_count > 0 and user_count > 0:
        print(
            f"Database already has {player_count} players, {match_count} matches and {user_count} users, skipping seed"
        )
        return

    workbook_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "FantasyCricket.xlsx")
    seeded_from_workbook = False
    if openpyxl is None:
        print("openpyxl not installed, skipping workbook seed")
    elif os.path.exists(workbook_path):
        print("Seeding database from FantasyCricket.xlsx")
        wb = openpyxl.load_workbook(workbook_path, read_only=True)

        ws = wb["Players"]
        workbook_player_count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                break
            db.execute(
                """
                INSERT INTO players (id, name, team, role, aliases)
                VALUES (?, ?, ?, ?, ?)
                ON CONFLICT (id) DO UPDATE SET
                    name = ?,
                    team = ?,
                    role = ?,
                    aliases = ?
                """,
                (int(row[0]), row[1], row[2], row[3], row[4] or "", row[1], row[2], row[3], row[4] or ""),
            )
            workbook_player_count += 1

        ws = wb["Matches"]
        workbook_match_count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                break
            date_str = row[1].strftime("%Y-%m-%d")
            time_str = row[2].strftime("%H:%M")
            db.execute(
                """
                INSERT INTO matches (id, team1, team2, match_date, match_time, status, toss_time)
                VALUES (?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT (id) DO UPDATE SET
                    team1 = ?,
                    team2 = ?,
                    match_date = ?,
                    match_time = ?,
                    status = ?,
                    toss_time = ?
                """,
                (
                    int(row[0]),
                    row[3],
                    row[4],
                    date_str,
                    time_str,
                    "future",
                    compute_toss_time(date_str, time_str),
                    row[3],
                    row[4],
                    date_str,
                    time_str,
                    "future",
                    compute_toss_time(date_str, time_str),
                ),
            )
            workbook_match_count += 1

        db.commit()
        wb.close()
        seeded_from_workbook = workbook_player_count > 0 or workbook_match_count > 0
        print(f"Seeded: {workbook_player_count} players, {workbook_match_count} matches")
    else:
        print(f"Seed file not found at {workbook_path}, using fallback local seed")

    if user_count == 0:
        db.execute(
            """
            INSERT INTO users (firebase_uid, email, name, mobile, role, is_active)
            VALUES (?, ?, ?, ?, ?, ?)
            ON CONFLICT(firebase_uid) DO UPDATE SET
                email = excluded.email,
                name = excluded.name,
                mobile = excluded.mobile,
                role = excluded.role,
                is_active = excluded.is_active
            """,
            ("dev_local_admin", "local.admin@example.com", "Local Admin", "", "admin", 1),
        )

    if player_count == 0 and not seeded_from_workbook:
        fallback_players = [
            (900001, "Local Batter", "CSK", "BAT", ""),
            (900002, "Local Wicketkeeper", "MI", "WK", ""),
            (900003, "Local Pacer All-Rounder", "RCB", "AR", "", "p"),
            (900004, "Local Spinner Bowler", "KKR", "BOWL", "", "s"),
        ]
        for player in fallback_players:
            if len(player) == 5:
                pid, name, team, role, aliases = player
                player_type = None
            else:
                pid, name, team, role, aliases, player_type = player
            db.execute(
                """
                INSERT INTO players (id, name, team, role, aliases, type)
                VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT (id) DO UPDATE SET
                    name = excluded.name,
                    team = excluded.team,
                    role = excluded.role,
                    aliases = excluded.aliases,
                    type = excluded.type
                """,
                (pid, name, team, role, aliases, player_type),
            )

    if match_count == 0 and not seeded_from_workbook:
        match_date = (get_current_datetime() + timedelta(days=1)).strftime("%Y-%m-%d")
        match_time = "19:30"
        db.execute(
            """
            INSERT INTO matches (id, team1, team2, match_date, match_time, status, venue, toss_time)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(id) DO UPDATE SET
                team1 = excluded.team1,
                team2 = excluded.team2,
                match_date = excluded.match_date,
                match_time = excluded.match_time,
                status = excluded.status,
                venue = excluded.venue,
                toss_time = excluded.toss_time
            """,
            (
                900001,
                "CSK",
                "MI",
                match_date,
                match_time,
                "future",
                "M.A. Chidambaram Stadium",
                compute_toss_time(match_date, match_time),
            ),
        )

    db.commit()
    print("Fallback local seed applied" if not seeded_from_workbook else "Local seed completed")


def normalize_player_types():
    db = get_db()
    updated_rows = 0
    type_overrides = {
        "Prince Yadav": "p",
        "Ravi Bishnoi": "s",
    }

    for player_name, player_type in type_overrides.items():
        result = db.execute(
            "UPDATE players SET type = ? WHERE LOWER(name) = LOWER(?)",
            (player_type, player_name),
        )
        updated_rows += result.rowcount or 0

    db.commit()
    if updated_rows:
        print(f"Updated player types for {updated_rows} rows")


def _load_todays_live_matches():
    today_key = get_current_date_key()
    matches_data = data_service.get_cached_data("matches")
    prepared_matches = []

    for match in matches_data:
        status, _ = matches.compute_runtime_match_status(
            match["Date"],
            match["Time"],
            match.get("Status"),
        )
        match_copy = dict(match)
        match_copy["status"] = status
        prepared_matches.append(match_copy)

    today_live_matches = [
        match
        for match in prepared_matches
        if (
            (match.get("Date") == today_key and match.get("status") not in {"completed", "nr"})
            or match.get("status") in {"lineups", "live"}
        )
    ]
    live_match_ids = [int(match["MatchID"]) for match in today_live_matches]
    return prepared_matches, today_live_matches, live_match_ids


def _run_deferred_match_warmup():
    try:
        time.sleep(2)
        print("[BOOT] Deferred match warmup starting")

        try:
            from backend.services.weekend_tournament_service import detect_and_create_tournaments
            created = detect_and_create_tournaments()
            if created:
                print(f"[BOOT] Weekend tournaments created: {created}")
            else:
                print("[BOOT] No new weekend tournaments detected")
        except Exception as exc:
            print(f"[BOOT] Weekend tournament detection failed: {exc}")

        try:
            from backend.services.scraper import populate_match_venues
            populate_match_venues(get_db())
        except Exception as exc:
            print(f"[BOOT] Venue backfill failed: {exc}")

        try:
            print("[BOOT] Deferred leaderboard warmup starting")
            summary = leaderboard.refresh_leaderboard_cache_once()
            print(
                "[BOOT] Deferred leaderboard warmup complete "
                f"leaderboard={summary['leaderboard']} points_table={summary['points_table']}"
            )
        except Exception as exc:
            print(f"[BOOT] Deferred leaderboard warmup failed: {exc}")

        print("[BOOT] Deferred match warmup complete")
    except Exception as exc:
        print(f"[BOOT] Deferred warmup error: {exc}")


def bootstrap_app():
    global bootstrap_ready, bootstrap_error
    try:
        init_db()
        print("Database initialized")
        seed_db_if_needed()
        normalize_player_types()
        init_firebase()
        bootstrap_ready = True
        bootstrap_error = None
        print("Fantasy Cricket API critical path ready")

        warmup_thread = threading.Thread(target=_run_background_warmup, daemon=True)
        warmup_thread.start()
    except Exception as exc:
        bootstrap_error = str(exc)
        print(f"Bootstrap error: {exc}")
        raise


def start_bootstrap_if_needed():
    global bootstrap_started
    with bootstrap_lock:
        if bootstrap_started:
            return
        bootstrap_started = True
        thread = threading.Thread(target=bootstrap_app, daemon=True)
        thread.start()


def _run_background_warmup():
    global bootstrap_warmup_complete, bootstrap_warmup_error

    while True:
        try:
            print("[BOOT] Background warmup starting")
            data_service.prime_static_cache()
            prediction_summary = data_service.prime_score_prediction_cache()
            print(f"[BOOT] Primed score prediction cache matches={len(prediction_summary)}")
            completed_rank_summary = matches.refresh_completed_match_rank_cache_once()
            print(
                "[BOOT] Primed completed match rank cache "
                f"matches={completed_rank_summary['matches']} ranked={completed_rank_summary['ranked']}"
            )
            contestant_summary = data_service.prime_contestant_cache()
            print(f"[BOOT] Primed contestant cache matches={len(contestant_summary)}")
            super_team_summary = super_team_service.prime_super_team_cache()
            print(
                "[BOOT] Primed super team cache "
                f"teams={super_team_summary['teams']} standings={super_team_summary['standings']}"
            )
            user_team_summary = data_service.prime_user_team_summary_cache()
            print(f"[BOOT] Primed user team summary cache users={len(user_team_summary)}")
            weekend_summary = prime_weekend_tournament_cache()
            print(
                "[BOOT] Primed weekend tournament cache "
                f"current={weekend_summary['current']} upcoming={weekend_summary['upcoming']} "
                f"history={weekend_summary['history']} by_id={weekend_summary['by_id']}"
            )

            _, today_live_matches, live_match_ids = _load_todays_live_matches()
            prime_today_venue_cache(today_live_matches)

            players_data = data_service.get_cached_data("players")
            teams_data = data_service.get_teams_for_matches(live_match_ids)

            tournament.initialize(players_data, today_live_matches, teams_data)
            tournament.start_lineup_cache_scheduler()
            tournament.start_toss_cache_scheduler()
            tournament.start_scheduler()

            print("[BOOT] Priming live scores cache")
            try:
                prime_summary = scores.refresh_scores_response_cache_once(match_statuses={"live", "nr"})
                print(
                    "[BOOT] Live scores cache primed "
                    f"matches={prime_summary['matches']} "
                    f"eligible={prime_summary['eligible']} "
                    f"refreshed={prime_summary['refreshed']} "
                    f"errors={prime_summary['errors']}"
                )
            except Exception as exc:
                print(f"[BOOT] Live scores cache prime failed: {exc}")

            print("[BOOT] Starting scores cache scheduler")
            scores.start_scores_cache_scheduler()
            print("[BOOT] Starting leaderboard cache scheduler")
            leaderboard.start_leaderboard_cache_scheduler()

            bootstrap_warmup_complete = True
            bootstrap_warmup_error = None
            print("Fantasy Cricket API initial warmup complete")

            deferred_thread = threading.Thread(target=_run_deferred_match_warmup, daemon=True)
            deferred_thread.start()
            return
        except Exception as exc:
            bootstrap_warmup_error = str(exc)
            print(f"[BOOT] Background warmup error: {exc}")
            time.sleep(30)


@app.on_event("startup")
def startup():
    admin.set_tournament(tournament)
    start_bootstrap_if_needed()
    # Populate venue data from Cricbuzz in background (fails silently)
    import threading
    def _populate_venues():
        try:
            from backend.services.scraper import populate_match_venues
            from backend.database import get_db
            populate_match_venues(get_db())
        except Exception:
            pass
    threading.Thread(target=_populate_venues, daemon=True).start()


@app.get("/api/health")
def health():
    status = "ok" if bootstrap_ready else "starting"
    if bootstrap_ready and not bootstrap_warmup_complete:
        status = "warming"
    payload = {
        "status": status,
        "critical_ready": bootstrap_ready,
        "warmup_complete": bootstrap_warmup_complete,
    }
    if bootstrap_error:
        payload["bootstrap_error"] = bootstrap_error
    if bootstrap_warmup_error:
        payload["bootstrap_warmup_error"] = bootstrap_warmup_error
    return payload


@app.get("/api/status")
def status():
    if not bootstrap_ready:
        return {
            "status": "starting",
            "scheduler": "booting",
            "bootstrap_error": bootstrap_error,
        }

    from backend.database import get_db
    db = get_db()
    try:
        players = db.execute("SELECT COUNT(*) as cnt FROM players").fetchone()
        matches = db.execute("SELECT COUNT(*) as cnt FROM matches").fetchone()
        users = db.execute("SELECT COUNT(*) as cnt FROM users").fetchone()
        return {
            "status": "ok" if bootstrap_warmup_complete else "warming",
            "players": players["cnt"] if isinstance(players, dict) else players[0],
            "matches": matches["cnt"] if isinstance(matches, dict) else matches[0],
            "users": users["cnt"] if isinstance(users, dict) else users[0],
            "scheduler": "running",
            "warmup_complete": bootstrap_warmup_complete,
        }
    except Exception as e:
        return {"status": "error", "detail": str(e)}


# Serve React static files in production
STATIC_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "frontend", "dist")

if os.path.isdir(STATIC_DIR):
    class CacheControlledStaticFiles(StaticFiles):
        async def get_response(self, path: str, scope):
            response = await super().get_response(path, scope)
            if response.status_code == 200 and path:
                if path.endswith("index.html"):
                    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
                    response.headers["Pragma"] = "no-cache"
                    response.headers["Expires"] = "0"
                else:
                    response.headers["Cache-Control"] = "public, max-age=31536000, immutable"
            return response

    # Serve static assets (js, css, images)
    app.mount("/assets", CacheControlledStaticFiles(directory=os.path.join(STATIC_DIR, "assets")), name="static-assets")

    # Catch-all: serve static files if they exist, otherwise index.html for React Router
    @app.get("/{full_path:path}")
    async def serve_spa(request: Request, full_path: str):
        if full_path.startswith("api/"):
            return {"detail": "Not found"}

        # Check if it's a real file in dist/
        file_path = os.path.join(STATIC_DIR, full_path)
        if full_path and os.path.isfile(file_path):
            response = FileResponse(file_path)
            if full_path.endswith("index.html"):
                response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
                response.headers["Pragma"] = "no-cache"
                response.headers["Expires"] = "0"
            else:
                response.headers["Cache-Control"] = "public, max-age=31536000, immutable"
            return response

        # Otherwise serve index.html for React Router
        response = FileResponse(os.path.join(STATIC_DIR, "index.html"))
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"
        return response
